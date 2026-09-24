import csv
import io
import time

from openpyxl import load_workbook

from app import db


def _wait(client, job_id, predicate, timeout=8):
    deadline = time.time() + timeout
    last = None
    while time.time() < deadline:
        last = client.get(f"/api/jobs/{job_id}").json()
        if predicate(last):
            return last
        time.sleep(0.05)
    raise AssertionError(last)


def _rows(response):
    assert response.status_code == 200, response.text
    text = response.content.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


def _sheet(response):
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    sheet = book.active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(header, values)))
    return header, rows


def _matched_fixture(client):
    created = client.post(
        "/api/jobs",
        json={
            "mode": "fixture",
            "fixture_set": "pokemon",
            "max_pages": 2,
            "delay_sec": 0,
            "expand_related": False,
            "recheck_after_block": False,
        },
    )
    assert created.status_code == 201, created.text
    amazon = _wait(client, created.json()["id"], lambda item: item["status"] == "completed")
    match = client.post(f"/api/jobs/{amazon['id']}/tcg-match")
    assert match.status_code == 201, match.text
    _wait(client, match.json()["id"], lambda item: item["status"] == "completed")
    observations = client.get(f"/api/jobs/{amazon['id']}/observations").json()["items"]
    return amazon, {item["asin"]: item for item in observations}


def test_exports_cover_amazon_tcg_and_price_compare_without_inventing_prices(client):
    amazon, by_asin = _matched_fixture(client)
    job_id = amazon["id"]

    job_csv = _rows(client.get("/api/exports/amazon-products", params={"job_id": job_id, "format": "csv"}))
    assert len(job_csv) == 4
    booster = next(row for row in job_csv if row["ASIN"] == "B0PKMN0001")
    assert booster["Title"].startswith("Pokemon TCG")
    assert booster["Price"] == "143.99"
    assert booster["Bought last month"] == "1000"
    assert booster["Bought last month text"] == "1K+"
    assert booster["TCGPlayer price label"] == "Market"
    assert booster["TCGPlayer price"] == "139.99"
    assert booster["Product ID"]
    assert booster["Product URL"]
    sleeves = next(row for row in job_csv if row["ASIN"] == "B0PKMN0004")
    assert sleeves["TCGPlayer status"] == "unmatched"
    assert sleeves["TCGPlayer price"] == ""
    assert sleeves["TCGPlayer URL"] == ""

    filtered = _rows(
        client.get(
            "/api/exports/amazon-products",
            params={"job_id": job_id, "q": "Sleeves", "format": "csv"},
        )
    )
    assert [row["ASIN"] for row in filtered] == ["B0PKMN0004"]

    disposition = client.get(
        "/api/exports/amazon-products",
        params={"job_id": job_id, "format": "xlsx"},
    )
    assert "spreadsheetml" in disposition.headers["content-type"]
    assert f'filename="amazon-products-job-{job_id}-' in disposition.headers["content-disposition"]
    assert disposition.headers["content-disposition"].endswith('.xlsx"')
    header, sheet_rows = _sheet(disposition)
    assert "ASIN" in header
    assert len(sheet_rows) == 4
    assert any(row["ASIN"] == "B0PKMN0001" and row["Price"] == 143.99 for row in sheet_rows)

    catalog = client.get("/api/exports/amazon-products", params={"format": "csv", "q": "Charizard"})
    catalog_rows = _rows(catalog)
    assert [row["ASIN"] for row in catalog_rows] == ["B0PKMN0003"]
    assert "filename=\"amazon-products-" in catalog.headers["content-disposition"]
    assert f"job-{job_id}" not in catalog.headers["content-disposition"]
    assert _rows(client.get("/api/exports/amazon-products", params={"has_asin": "no"})) == []

    tcg = _rows(client.get("/api/exports/tcgplayer-matches", params={"job_id": job_id}))
    twilight = [row for row in tcg if row["ASIN"] == "B0PKMN0002"]
    assert len(twilight) == 2
    assert {row["Confirmed"] for row in twilight} == {"no"}
    assert sorted(float(row["Price"]) for row in twilight) == [44.95, 89.0]
    assert all(row["Match status"] == "needs_confirm" for row in twilight)
    assert all(row["TCGPlayer URL"].startswith("https://www.tcgplayer.com/") for row in twilight)
    booster_tcg = next(row for row in tcg if row["ASIN"] == "B0PKMN0001")
    assert booster_tcg["Confirmed"] == "yes"
    assert booster_tcg["Price label"] == "Market"
    assert "Low" in booster_tcg["Other prices"]
    sleeves_tcg = next(row for row in tcg if row["ASIN"] == "B0PKMN0004")
    assert sleeves_tcg["Match status"] == "unmatched"
    assert sleeves_tcg["Price"] == ""
    assert sleeves_tcg["TCGPlayer name"] == ""

    compared = _rows(client.get("/api/exports/price-compare", params={"job_id": job_id, "format": "csv"}))
    asins = {row["ASIN"] for row in compared}
    assert "B0PKMN0001" in asins
    assert "B0PKMN0003" in asins
    assert "B0PKMN0002" not in asins
    assert "B0PKMN0004" not in asins
    booster_gap = next(row for row in compared if row["ASIN"] == "B0PKMN0001")
    assert booster_gap["Amazon price"] == "143.99"
    assert booster_gap["TCGPlayer price"] == "139.99"
    assert booster_gap["Difference"] == "4.0"
    assert booster_gap["Lower"] == "tcgplayer"
    assert booster_gap["Bought last month text"] == "1K+"
    tin = next(row for row in compared if row["ASIN"] == "B0PKMN0003")
    assert float(tin["Difference"]) == -1.0
    assert tin["Lower"] == "amazon"

    narrowed = _rows(
        client.get("/api/exports/price-compare", params={"job_id": job_id, "q": "Charizard", "format": "csv"})
    )
    assert [row["ASIN"] for row in narrowed] == ["B0PKMN0003"]

    etb = by_asin["B0PKMN0002"]
    blocked = client.get("/api/exports/price-compare", params={"product_id": etb["product_id"]})
    assert blocked.status_code == 409
    product = client.get(f"/api/products/{etb['product_id']}").json()
    chosen = next(item for item in product["tcg_candidates"] if item["name"] == "Twilight Masquerade Elite Trainer Box")
    confirmed = client.post(
        f"/api/products/{etb['product_id']}/tcg-confirm",
        json={"candidate_id": chosen["id"]},
    )
    assert confirmed.status_code == 200, confirmed.text
    after = _rows(client.get("/api/exports/price-compare", params={"product_id": etb["product_id"], "format": "csv"}))
    assert len(after) == 1
    assert after[0]["TCGPlayer name"] == "Twilight Masquerade Elite Trainer Box"
    assert after[0]["Amazon price"] == "49.5"
    assert after[0]["TCGPlayer price"] == "44.95"
    assert after[0]["Difference"] == "4.55"
    assert after[0]["Lower"] == "tcgplayer"
    tcg_after = _rows(client.get("/api/exports/tcgplayer-matches", params={"product_id": etb["product_id"]}))
    assert sum(row["Confirmed"] == "yes" for row in tcg_after) == 1

    with db._tx() as conn:
        conn.execute(
            "UPDATE scrape_observations SET price = NULL WHERE product_id = ?",
            (by_asin["B0PKMN0001"]["product_id"],),
        )
    blank = _rows(
        client.get(
            "/api/exports/price-compare",
            params={"product_id": by_asin["B0PKMN0001"]["product_id"]},
        )
    )
    assert blank[0]["Amazon price"] == ""
    assert blank[0]["TCGPlayer price"] == "139.99"
    assert blank[0]["Difference"] == ""
    assert blank[0]["Lower"] == ""

    listed = client.get("/api/exports/tcgplayer-matches", params={"format": "json", "q": "Twilight"}).json()
    assert listed["total"] == 2
    assert {item["price"] for item in listed["items"]} == {44.95, 89.0}

    assert client.get("/api/exports/not-a-list").status_code == 404
    assert client.get("/api/exports/amazon-products", params={"format": "pdf"}).status_code == 400
    assert client.get("/api/exports/amazon-products", params={"job_id": 999999}).status_code == 404

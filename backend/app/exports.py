"""CSV and Excel downloads for the Amazon, TCGPlayer, and price-compare lists."""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.db import _escape_like, _tx, list_observations

AMAZON_JOB_COLUMNS = [
    ("product_id", "Product ID"),
    ("observation_id", "Observation ID"),
    ("job_id", "Job ID"),
    ("title", "Title"),
    ("asin", "ASIN"),
    ("price", "Price"),
    ("currency", "Currency"),
    ("list_price", "List price"),
    ("rating", "Rating"),
    ("review_count", "Reviews"),
    ("bought_past_month", "Bought last month"),
    ("bought_past_month_text", "Bought last month text"),
    ("page_number", "Page"),
    ("source", "Source"),
    ("image_url", "Image URL"),
    ("product_url", "Product URL"),
    ("seller", "Seller"),
    ("observed_at", "Seen at"),
    ("availability_snippet", "Availability"),
    ("badges", "Badges"),
    ("tcg_status", "TCGPlayer status"),
    ("tcg_name", "TCGPlayer name"),
    ("tcg_set", "TCGPlayer set"),
    ("tcg_price_label", "TCGPlayer price label"),
    ("tcg_price", "TCGPlayer price"),
    ("tcg_currency", "TCGPlayer currency"),
    ("tcg_url", "TCGPlayer URL"),
]

AMAZON_CATALOG_COLUMNS = [
    ("product_id", "Product ID"),
    ("title", "Title"),
    ("asin", "ASIN"),
    ("image_url", "Image URL"),
    ("product_url", "Product URL"),
    ("category_breadcrumbs", "Breadcrumbs"),
    ("first_seen_at", "First seen"),
    ("last_seen_at", "Last seen"),
    ("observation_count", "Snapshots"),
    ("latest_price", "Latest price"),
    ("currency", "Currency"),
    ("list_price", "List price"),
    ("rating", "Rating"),
    ("review_count", "Reviews"),
    ("bought_past_month", "Bought last month"),
    ("bought_past_month_text", "Bought last month text"),
    ("tcg_status", "TCGPlayer status"),
    ("tcg_name", "TCGPlayer name"),
    ("tcg_set", "TCGPlayer set"),
    ("tcg_price_label", "TCGPlayer price label"),
    ("tcg_price", "TCGPlayer price"),
    ("tcg_currency", "TCGPlayer currency"),
    ("tcg_url", "TCGPlayer URL"),
]

TCG_COLUMNS = [
    ("product_id", "Product ID"),
    ("asin", "ASIN"),
    ("amazon_title", "Amazon title"),
    ("amazon_url", "Amazon URL"),
    ("match_status", "Match status"),
    ("query", "Search query"),
    ("candidate_id", "Candidate ID"),
    ("confirmed", "Confirmed"),
    ("tcg_name", "TCGPlayer name"),
    ("set_name", "Set"),
    ("tcg_url", "TCGPlayer URL"),
    ("image_url", "Image URL"),
    ("price_label", "Price label"),
    ("price", "Price"),
    ("currency", "Currency"),
    ("other_prices", "Other prices"),
    ("confidence", "Confidence"),
]

COMPARE_COLUMNS = [
    ("product_id", "Product ID"),
    ("amazon_title", "Amazon title"),
    ("asin", "ASIN"),
    ("amazon_price", "Amazon price"),
    ("amazon_currency", "Amazon currency"),
    ("bought_past_month", "Bought last month"),
    ("bought_past_month_text", "Bought last month text"),
    ("amazon_url", "Amazon URL"),
    ("amazon_image_url", "Amazon image URL"),
    ("tcg_name", "TCGPlayer name"),
    ("tcg_set", "Set"),
    ("tcg_price_label", "TCGPlayer price label"),
    ("tcg_price", "TCGPlayer price"),
    ("tcg_currency", "TCGPlayer currency"),
    ("tcg_other_prices", "TCGPlayer other prices"),
    ("tcg_url", "TCGPlayer URL"),
    ("tcg_image_url", "TCGPlayer image URL"),
    ("difference", "Difference"),
    ("lower", "Lower"),
]

DATASETS = {
    "amazon-products": "amazon-products",
    "tcgplayer-matches": "tcgplayer-matches",
    "price-compare": "price-compare",
}


def export_filename(kind: str, extension: str, *, job_id: int | None) -> str:
    day = datetime.now(timezone.utc).strftime("%Y%m%d")
    stem = DATASETS[kind]
    if job_id is not None:
        return f"{stem}-job-{job_id}-{day}.{extension}"
    return f"{stem}-{day}.{extension}"


def render_csv(columns: list[tuple[str, str]], rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([label for _key, label in columns])
    for row in rows:
        writer.writerow([_csv_cell(row.get(key)) for key, _label in columns])
    return buffer.getvalue().encode("utf-8-sig")


def render_xlsx(columns: list[tuple[str, str]], rows: list[dict], *, sheet_name: str) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name[:31]
    sheet.append([label for _key, label in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_xlsx_cell(row.get(key)) for key, _label in columns])
    for index, (_key, label) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(label) + 2, 14), 42)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    payload = io.BytesIO()
    book.save(payload)
    return payload.getvalue()


def amazon_rows(
    *,
    job_id: int | None,
    product_id: int | None,
    q: str | None,
    min_price: float | None,
    max_price: float | None,
    min_rating: float | None,
    min_bought: int | None,
    sort: str,
    has_asin: str | None,
    last_seen_after: str | None,
    last_seen_before: str | None,
) -> tuple[list[tuple[str, str]], list[dict], str]:
    if job_id is not None:
        page = list_observations(
            job_id,
            q=q,
            min_price=min_price,
            max_price=max_price,
            min_rating=min_rating,
            min_bought=min_bought,
            sort=sort,
            limit=10_000_000,
            offset=0,
        )
        rows = []
        for item in page["items"]:
            if product_id is not None and int(item["product_id"]) != product_id:
                continue
            badges = item.get("badges") or []
            rows.append(
                {
                    "product_id": item["product_id"],
                    "observation_id": item["observation_id"],
                    "job_id": item["job_id"],
                    "title": item["title"],
                    "asin": item.get("asin"),
                    "price": item.get("price"),
                    "currency": item.get("currency"),
                    "list_price": item.get("list_price"),
                    "rating": item.get("rating"),
                    "review_count": item.get("review_count"),
                    "bought_past_month": item.get("bought_past_month"),
                    "bought_past_month_text": item.get("bought_past_month_text"),
                    "page_number": item.get("page_number"),
                    "source": item.get("source"),
                    "image_url": item.get("image_url"),
                    "product_url": _amazon_url(item.get("product_url"), item.get("asin")),
                    "seller": item.get("seller"),
                    "observed_at": item.get("observed_at"),
                    "availability_snippet": item.get("availability_snippet"),
                    "badges": " | ".join(badges) if badges else None,
                    "tcg_status": item.get("tcg_status"),
                    "tcg_name": item.get("tcg_name"),
                    "tcg_set": item.get("tcg_set"),
                    "tcg_price_label": item.get("tcg_price_label"),
                    "tcg_price": item.get("tcg_price"),
                    "tcg_currency": item.get("tcg_currency"),
                    "tcg_url": item.get("tcg_url"),
                }
            )
        return AMAZON_JOB_COLUMNS, rows, "Amazon products"
    rows = _catalog_amazon_rows(
        q=q,
        has_asin=has_asin,
        last_seen_after=last_seen_after,
        last_seen_before=last_seen_before,
        product_id=product_id,
    )
    return AMAZON_CATALOG_COLUMNS, rows, "Amazon products"


def tcg_rows(
    *,
    job_id: int | None,
    product_id: int | None,
    q: str | None,
    min_price: float | None,
    max_price: float | None,
    min_rating: float | None,
    min_bought: int | None,
    has_asin: str | None,
    last_seen_after: str | None,
    last_seen_before: str | None,
) -> list[dict]:
    where, args = _scope_where(
        job_id=job_id,
        product_id=product_id,
        q=q,
        min_price=min_price,
        max_price=max_price,
        min_rating=min_rating,
        min_bought=min_bought,
        has_asin=has_asin,
        last_seen_after=last_seen_after,
        last_seen_before=last_seen_before,
        include_tcg_text=job_id is None,
    )
    clause = " AND ".join(where)
    with _tx() as conn:
        fetched = conn.execute(
            f"""
            SELECT p.id AS product_id, p.asin, p.title AS amazon_title, p.product_url,
                   m.status AS match_status, m.query AS query, m.tcg_url AS confirmed_url,
                   m.tcg_name AS confirmed_name, m.tcg_set AS confirmed_set,
                   m.image_url AS confirmed_image, m.price AS confirmed_price,
                   m.price_label AS confirmed_price_label, m.currency AS confirmed_currency,
                   m.confidence AS confirmed_confidence,
                   c.id AS candidate_id, c.name AS candidate_name, c.set_name AS candidate_set,
                   c.url AS candidate_url, c.image_url AS candidate_image,
                   c.price AS candidate_price, c.price_label AS candidate_price_label,
                   c.currency AS candidate_currency, c.prices_json AS candidate_prices,
                   c.confidence AS candidate_confidence
            FROM products p
            JOIN tcgplayer_matches m ON m.product_id = p.id
            LEFT JOIN tcgplayer_candidates c ON c.product_id = p.id
            WHERE {clause}
            ORDER BY p.title COLLATE NOCASE, p.id, c.position, c.id
            """,
            args,
        ).fetchall()
    grouped: dict[int, list] = {}
    order: list[int] = []
    for row in fetched:
        pid = int(row["product_id"])
        if pid not in grouped:
            grouped[pid] = []
            order.append(pid)
        grouped[pid].append(row)
    exported: list[dict] = []
    for pid in order:
        group = grouped[pid]
        base = group[0]
        candidates = [row for row in group if row["candidate_id"] is not None]
        amazon_url = _amazon_url(base["product_url"], base["asin"])
        if not candidates:
            exported.append(_tcg_match_row(base, amazon_url))
            continue
        confirmed_url = base["confirmed_url"] if base["match_status"] == "matched" else None
        saw_confirmed = False
        for candidate in candidates:
            is_confirmed = bool(confirmed_url and candidate["candidate_url"] == confirmed_url)
            saw_confirmed = saw_confirmed or is_confirmed
            exported.append(
                {
                    "product_id": pid,
                    "asin": base["asin"],
                    "amazon_title": base["amazon_title"],
                    "amazon_url": amazon_url,
                    "match_status": base["match_status"],
                    "query": base["query"],
                    "candidate_id": int(candidate["candidate_id"]),
                    "confirmed": "yes" if is_confirmed else "no",
                    "tcg_name": candidate["candidate_name"],
                    "set_name": candidate["candidate_set"],
                    "tcg_url": candidate["candidate_url"],
                    "image_url": candidate["candidate_image"],
                    "price_label": candidate["candidate_price_label"],
                    "price": candidate["candidate_price"],
                    "currency": candidate["candidate_currency"],
                    "other_prices": _other_prices(candidate["candidate_prices"]),
                    "confidence": candidate["candidate_confidence"],
                }
            )
        if confirmed_url and not saw_confirmed:
            exported.append(_tcg_match_row(base, amazon_url))
    return exported


def price_compare_rows(
    *,
    job_id: int | None,
    product_id: int | None,
    q: str | None,
    min_price: float | None,
    max_price: float | None,
    min_rating: float | None,
    min_bought: int | None,
    has_asin: str | None,
    last_seen_after: str | None,
    last_seen_before: str | None,
) -> list[dict]:
    where, args = _scope_where(
        job_id=job_id,
        product_id=product_id,
        q=q,
        min_price=min_price,
        max_price=max_price,
        min_rating=min_rating,
        min_bought=min_bought,
        has_asin=has_asin,
        last_seen_after=last_seen_after,
        last_seen_before=last_seen_before,
        include_tcg_text=job_id is None,
    )
    where.append("m.status = 'matched'")
    clause = " AND ".join(where)
    with _tx() as conn:
        fetched = conn.execute(
            f"""
            SELECT p.id AS product_id, p.title AS amazon_title, p.asin, p.image_url AS amazon_image_url,
                   p.product_url,
                   obs.price AS amazon_price, obs.currency AS amazon_currency,
                   obs.bought_past_month, obs.bought_past_month_text,
                   m.tcg_name, m.tcg_set, m.price_label AS tcg_price_label, m.price AS tcg_price,
                   m.currency AS tcg_currency, m.tcg_url, m.image_url AS tcg_image_url,
                   (
                     SELECT c.prices_json
                     FROM tcgplayer_candidates c
                     WHERE c.product_id = p.id AND c.url = m.tcg_url
                     ORDER BY c.position, c.id
                     LIMIT 1
                   ) AS prices_json
            FROM products p
            JOIN tcgplayer_matches m ON m.product_id = p.id
            LEFT JOIN scrape_observations obs ON obs.id = (
                SELECT o.id
                FROM scrape_observations o
                WHERE o.product_id = p.id
                ORDER BY o.observed_at DESC, o.id DESC
                LIMIT 1
            )
            WHERE {clause}
            ORDER BY p.title COLLATE NOCASE, p.id
            """,
            args,
        ).fetchall()
    rows = []
    for row in fetched:
        amazon_price = row["amazon_price"]
        tcg_price = row["tcg_price"]
        difference = None
        lower = None
        if amazon_price is not None and tcg_price is not None:
            difference = round(float(amazon_price) - float(tcg_price), 2)
            if abs(difference) < 0.005:
                lower = "same"
            elif difference > 0:
                lower = "tcgplayer"
            else:
                lower = "amazon"
        rows.append(
            {
                "product_id": int(row["product_id"]),
                "amazon_title": row["amazon_title"],
                "asin": row["asin"],
                "amazon_price": amazon_price,
                "amazon_currency": row["amazon_currency"],
                "bought_past_month": row["bought_past_month"],
                "bought_past_month_text": row["bought_past_month_text"],
                "amazon_url": _amazon_url(row["product_url"], row["asin"]),
                "amazon_image_url": row["amazon_image_url"],
                "tcg_name": row["tcg_name"],
                "tcg_set": row["tcg_set"],
                "tcg_price_label": row["tcg_price_label"],
                "tcg_price": tcg_price,
                "tcg_currency": row["tcg_currency"],
                "tcg_other_prices": _other_prices(row["prices_json"]),
                "tcg_url": row["tcg_url"],
                "tcg_image_url": row["tcg_image_url"],
                "difference": difference,
                "lower": lower,
            }
        )
    return rows


def product_match_state(product_id: int) -> str | None:
    """Return the product's match status, or None when the product does not exist."""
    with _tx() as conn:
        product = conn.execute("SELECT id FROM products WHERE id = ?", (product_id,)).fetchone()
        if product is None:
            return None
        match = conn.execute(
            "SELECT status FROM tcgplayer_matches WHERE product_id = ?",
            (product_id,),
        ).fetchone()
    if match is None:
        return ""
    return str(match["status"])


def _catalog_amazon_rows(
    *,
    q: str | None,
    has_asin: str | None,
    last_seen_after: str | None,
    last_seen_before: str | None,
    product_id: int | None,
) -> list[dict]:
    where = ["1 = 1"]
    args: list = []
    if product_id is not None:
        where.append("p.id = ?")
        args.append(product_id)
    if q:
        like = f"%{_escape_like(q)}%"
        where.append(
            "(p.title LIKE ? ESCAPE '\\' OR IFNULL(p.asin, '') LIKE ? ESCAPE '\\' "
            "OR IFNULL(p.product_url, '') LIKE ? ESCAPE '\\')"
        )
        args.extend([like, like, like])
    if has_asin == "yes":
        where.append("p.asin IS NOT NULL")
    elif has_asin == "no":
        where.append("p.asin IS NULL")
    if last_seen_after:
        where.append("p.last_seen_at >= ?")
        args.append(last_seen_after)
    if last_seen_before:
        where.append("p.last_seen_at < ?")
        args.append(last_seen_before)
    clause = " AND ".join(where)
    with _tx() as conn:
        fetched = conn.execute(
            f"""
            SELECT p.id AS product_id, p.title, p.asin, p.image_url, p.product_url,
                   p.category_breadcrumbs, p.first_seen_at, p.last_seen_at,
                   (SELECT COUNT(*) FROM scrape_observations o WHERE o.product_id = p.id)
                     AS observation_count,
                   obs.price AS latest_price, obs.currency, obs.list_price, obs.rating,
                   obs.review_count, obs.bought_past_month, obs.bought_past_month_text,
                   tcg.status AS tcg_status, tcg.tcg_name, tcg.tcg_set,
                   tcg.price_label AS tcg_price_label, tcg.price AS tcg_price,
                   tcg.currency AS tcg_currency, tcg.tcg_url
            FROM products p
            LEFT JOIN tcgplayer_matches tcg ON tcg.product_id = p.id
            LEFT JOIN scrape_observations obs ON obs.id = (
                SELECT o.id FROM scrape_observations o
                WHERE o.product_id = p.id
                ORDER BY o.observed_at DESC, o.id DESC
                LIMIT 1
            )
            WHERE {clause}
            ORDER BY p.last_seen_at DESC, p.id DESC
            """,
            args,
        ).fetchall()
    rows = []
    for row in fetched:
        item = dict(row)
        item["product_url"] = _amazon_url(item.get("product_url"), item.get("asin"))
        rows.append(item)
    return rows


def _scope_where(
    *,
    job_id: int | None,
    product_id: int | None,
    q: str | None,
    min_price: float | None,
    max_price: float | None,
    min_rating: float | None,
    min_bought: int | None,
    has_asin: str | None,
    last_seen_after: str | None,
    last_seen_before: str | None,
    include_tcg_text: bool,
) -> tuple[list[str], list]:
    where = ["1 = 1"]
    args: list = []
    if product_id is not None:
        where.append("p.id = ?")
        args.append(product_id)
    if job_id is not None:
        obs_where = ["o.job_id = ?"]
        obs_args: list = [job_id]
        if q:
            like = f"%{_escape_like(q)}%"
            obs_where.append(
                "(fp.title LIKE ? ESCAPE '\\' OR IFNULL(fp.asin, '') LIKE ? ESCAPE '\\')"
            )
            obs_args.extend([like, like])
        if min_price is not None:
            obs_where.append("o.price >= ?")
            obs_args.append(min_price)
        if max_price is not None:
            obs_where.append("o.price <= ?")
            obs_args.append(max_price)
        if min_rating is not None:
            obs_where.append("o.rating >= ?")
            obs_args.append(min_rating)
        if min_bought is not None:
            obs_where.append("o.bought_past_month >= ?")
            obs_args.append(min_bought)
        where.append(
            "p.id IN (SELECT o.product_id FROM scrape_observations o "
            "JOIN products fp ON fp.id = o.product_id WHERE "
            + " AND ".join(obs_where)
            + ")"
        )
        args.extend(obs_args)
        return where, args
    if q:
        like = f"%{_escape_like(q)}%"
        pieces = [
            "p.title LIKE ? ESCAPE '\\'",
            "IFNULL(p.asin, '') LIKE ? ESCAPE '\\'",
            "IFNULL(p.product_url, '') LIKE ? ESCAPE '\\'",
        ]
        args.extend([like, like, like])
        if include_tcg_text:
            pieces.append("IFNULL(m.tcg_name, '') LIKE ? ESCAPE '\\'")
            pieces.append("IFNULL(m.tcg_set, '') LIKE ? ESCAPE '\\'")
            args.extend([like, like])
        where.append("(" + " OR ".join(pieces) + ")")
    if has_asin == "yes":
        where.append("p.asin IS NOT NULL")
    elif has_asin == "no":
        where.append("p.asin IS NULL")
    if last_seen_after:
        where.append("p.last_seen_at >= ?")
        args.append(last_seen_after)
    if last_seen_before:
        where.append("p.last_seen_at < ?")
        args.append(last_seen_before)
    return where, args


def _tcg_match_row(base, amazon_url: str | None) -> dict:
    confirmed = ""
    if base["match_status"] == "matched" and base["confirmed_url"]:
        confirmed = "yes"
    elif base["match_status"] == "needs_confirm":
        confirmed = "no"
    return {
        "product_id": int(base["product_id"]),
        "asin": base["asin"],
        "amazon_title": base["amazon_title"],
        "amazon_url": amazon_url,
        "match_status": base["match_status"],
        "query": base["query"],
        "candidate_id": None,
        "confirmed": confirmed,
        "tcg_name": base["confirmed_name"],
        "set_name": base["confirmed_set"],
        "tcg_url": base["confirmed_url"],
        "image_url": base["confirmed_image"],
        "price_label": base["confirmed_price_label"],
        "price": base["confirmed_price"],
        "currency": base["confirmed_currency"],
        "other_prices": None,
        "confidence": base["confirmed_confidence"],
    }


def _other_prices(raw) -> str | None:
    if raw is None or raw == "":
        return None
    try:
        prices = json.loads(raw) if isinstance(raw, str) else raw
    except json.JSONDecodeError:
        return None
    if not isinstance(prices, list):
        return None
    parts: list[str] = []
    for item in prices:
        if not isinstance(item, dict) or item.get("amount") is None:
            continue
        label = item.get("label") or "Price"
        parts.append(f"{label} {item['amount']}")
    return "; ".join(parts) or None


def _amazon_url(product_url: str | None, asin: str | None) -> str | None:
    if product_url:
        return product_url
    if asin:
        return f"https://www.amazon.com/dp/{asin}"
    return None


def _csv_cell(value):
    if value is None:
        return ""
    return value


def _xlsx_cell(value):
    if value is None:
        return None
    return value
